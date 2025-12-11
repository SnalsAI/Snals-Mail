"""
Attachment Text Extractor Service.

Estrae testo da vari tipi di allegati (PDF, DOCX, TXT, ecc.)
per passarlo all'LLM durante categorizzazione e interpretazione.
"""
import os
import logging
from pathlib import Path
from typing import Optional, List, Dict

logger = logging.getLogger(__name__)


class AttachmentTextExtractor:
    """Servizio per estrarre testo dagli allegati."""

    # Dimensione massima per estrazione testo (5 MB)
    MAX_FILE_SIZE = 5 * 1024 * 1024

    # Numero massimo caratteri da estrarre per file
    MAX_TEXT_LENGTH = 10000

    def __init__(self):
        """Inizializza l'estrattore."""
        self.extractors = {
            '.txt': self._extract_text_plain,
            '.pdf': self._extract_text_pdf,
            '.docx': self._extract_text_docx,
            '.doc': self._extract_text_doc,
            '.eml': self._extract_text_eml,
            '.msg': self._extract_text_msg,
        }

    def extract_from_attachments(self, attachment_paths: List[str]) -> str:
        """
        Estrae testo da una lista di allegati.

        Args:
            attachment_paths: Lista di path degli allegati

        Returns:
            Testo estratto da tutti gli allegati, concatenato
        """
        if not attachment_paths:
            return ""

        extracted_texts = []

        for path in attachment_paths:
            # Scarta file tecnici PEC (.p7s firma digitale, .xml daticert)
            ext = Path(path).suffix.lower()
            if ext in ['.p7s', '.xml']:
                logger.debug(f"🗑️ Scartato file tecnico PEC: {os.path.basename(path)}")
                continue

            try:
                text = self.extract_text(path)
                if text:
                    filename = os.path.basename(path)
                    extracted_texts.append(f"\n--- ALLEGATO: {filename} ---\n{text}\n")
            except Exception as e:
                logger.warning(f"Errore estrazione testo da {path}: {e}")
                continue

        return "\n".join(extracted_texts)

    def extract_text(self, file_path: str) -> Optional[str]:
        """
        Estrae testo da un singolo file.

        Args:
            file_path: Path del file

        Returns:
            Testo estratto o None
        """
        if not os.path.exists(file_path):
            logger.warning(f"File non trovato: {file_path}")
            return None

        # Verifica dimensione file
        file_size = os.path.getsize(file_path)
        if file_size > self.MAX_FILE_SIZE:
            logger.warning(f"File troppo grande ({file_size} bytes): {file_path}")
            return f"[File troppo grande per estrazione: {file_size / 1024 / 1024:.1f} MB]"

        # Determina estensione
        ext = Path(file_path).suffix.lower()

        # Usa estrattore specifico
        extractor = self.extractors.get(ext)
        if extractor:
            try:
                text = extractor(file_path)
                if text:
                    # Limita lunghezza
                    if len(text) > self.MAX_TEXT_LENGTH:
                        text = text[:self.MAX_TEXT_LENGTH] + "\n[... testo troncato ...]"
                    return text
            except Exception as e:
                logger.error(f"Errore estrazione da {file_path}: {e}")
                return f"[Errore estrazione testo: {str(e)}]"
        else:
            logger.debug(f"Nessun estrattore per estensione {ext}: {file_path}")
            return f"[Tipo file non supportato: {ext}]"

        return None

    def _extract_text_plain(self, file_path: str) -> str:
        """Estrae testo da file TXT."""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    def _extract_text_pdf(self, file_path: str) -> str:
        """
        Estrae testo da file PDF.

        Usa pdfplumber come metodo principale (migliore per layout complessi, tabelle),
        con fallback a PyPDF2 se pdfplumber non è disponibile.
        """
        # Prova prima con pdfplumber (migliore per PDF complessi)
        try:
            import pdfplumber
            text_parts = []

            with pdfplumber.open(file_path) as pdf:
                num_pages = len(pdf.pages)
                max_pages = min(num_pages, 10)  # Limita a prime 10 pagine

                for page_num in range(max_pages):
                    page = pdf.pages[page_num]
                    page_text = page.extract_text()

                    if page_text:
                        text_parts.append(page_text)

                    # Estrai anche tabelle se presenti
                    tables = page.extract_tables()
                    if tables:
                        for table in tables:
                            # Converti tabella in testo
                            table_text = "\n".join(["\t".join([str(cell) if cell else "" for cell in row]) for row in table])
                            if table_text.strip():
                                text_parts.append(f"\n[TABELLA]\n{table_text}\n")

                if num_pages > max_pages:
                    text_parts.append(f"\n[... {num_pages - max_pages} pagine aggiuntive non estratte ...]")

            extracted = "\n".join(text_parts)
            if extracted.strip():
                logger.debug(f"✅ PDF estratto con pdfplumber: {len(extracted)} caratteri")
                return extracted
            else:
                # Se pdfplumber non ha estratto nulla, prova PyPDF2
                logger.debug("pdfplumber non ha estratto testo, provo PyPDF2...")
                return self._extract_text_pdf_fallback(file_path)

        except ImportError:
            logger.debug("pdfplumber non disponibile, uso PyPDF2 come fallback")
            return self._extract_text_pdf_fallback(file_path)
        except Exception as e:
            logger.warning(f"Errore estrazione PDF con pdfplumber: {e}, provo PyPDF2...")
            return self._extract_text_pdf_fallback(file_path)

    def _extract_text_pdf_fallback(self, file_path: str) -> str:
        """Fallback per estrazione PDF usando PyPDF2."""
        try:
            import PyPDF2
            text_parts = []

            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                num_pages = len(pdf_reader.pages)
                max_pages = min(num_pages, 10)

                for page_num in range(max_pages):
                    page = pdf_reader.pages[page_num]
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)

                if num_pages > max_pages:
                    text_parts.append(f"\n[... {num_pages - max_pages} pagine aggiuntive non estratte ...]")

            extracted = "\n".join(text_parts)
            if extracted.strip():
                logger.debug(f"✅ PDF estratto con PyPDF2: {len(extracted)} caratteri")
                return extracted
            else:
                logger.warning(f"PDF {file_path} sembra vuoto o non estraibile")
                return "[PDF senza testo estraibile - potrebbe contenere solo immagini]"

        except ImportError:
            logger.error("Né pdfplumber né PyPDF2 sono installati!")
            return "[Nessuna libreria PDF disponibile - installare pdfplumber o PyPDF2]"
        except Exception as e:
            logger.error(f"Errore estrazione PDF con PyPDF2 {file_path}: {e}")
            return f"[Errore lettura PDF: {str(e)}]"

    def _extract_text_docx(self, file_path: str) -> str:
        """Estrae testo da file DOCX."""
        try:
            import docx

            doc = docx.Document(file_path)
            text_parts = []

            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)

            # Estrai anche dalle tabelle
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)

            return "\n".join(text_parts)

        except ImportError:
            logger.warning("python-docx non installato, impossibile estrarre testo da DOCX")
            return "[python-docx non disponibile - installare con: pip install python-docx]"
        except Exception as e:
            logger.error(f"Errore estrazione DOCX {file_path}: {e}")
            return f"[Errore lettura DOCX: {str(e)}]"

    def _extract_text_doc(self, file_path: str) -> str:
        """Estrae testo da file DOC (formato vecchio)."""
        try:
            import textract
            text = textract.process(file_path).decode('utf-8', errors='ignore')
            return text
        except ImportError:
            logger.warning("textract non installato, impossibile estrarre testo da DOC")
            return "[textract non disponibile - per file .doc considerare conversione a .docx]"
        except Exception as e:
            logger.error(f"Errore estrazione DOC {file_path}: {e}")
            return f"[Errore lettura DOC: {str(e)}]"

    def parse_eml_complete(self, file_path: str) -> Optional[Dict[str, str]]:
        """
        Parsa completamente un file EML estraendo mittente, oggetto e corpo.
        Usato per buste PEC.

        Args:
            file_path: Path del file EML

        Returns:
            Dict con 'from', 'subject', 'body' o None se errore
        """
        try:
            import email
            from email import policy
            from email.utils import parseaddr

            with open(file_path, 'rb') as f:
                msg = email.message_from_binary_file(f, policy=policy.default)

            # Estrai mittente
            from_header = msg.get('From', '')
            _, sender_email = parseaddr(from_header)

            # Estrai oggetto
            subject = msg.get('Subject', '')

            # Estrai corpo
            body = ""
            if msg.is_multipart():
                text_parts = []
                html_parts = []
                for part in msg.walk():
                    content_type = part.get_content_type()
                    if content_type == "text/plain":
                        try:
                            text_parts.append(part.get_content())
                        except:
                            pass
                    elif content_type == "text/html":
                        try:
                            html_parts.append(part.get_content())
                        except:
                            pass

                # Preferisci testo plain, fallback a HTML
                if text_parts:
                    body = "\n".join(text_parts)
                elif html_parts:
                    body = "\n".join(html_parts)
            else:
                try:
                    body = msg.get_content()
                except:
                    body = str(msg.get_payload())

            result = {
                'from': sender_email or from_header,
                'subject': subject,
                'body': body[:self.MAX_TEXT_LENGTH]  # Limita lunghezza
            }

            logger.info(f"✉️ EML parsato: From={result['from'][:50]}, Subject={result['subject'][:50]}")
            return result

        except Exception as e:
            logger.error(f"Errore parsing completo EML {file_path}: {e}")
            return None

    def _extract_text_eml(self, file_path: str) -> str:
        """Estrae testo da file EML (email salvate)."""
        try:
            import email
            from email import policy

            with open(file_path, 'rb') as f:
                msg = email.message_from_binary_file(f, policy=policy.default)

            # Estrai corpo
            if msg.is_multipart():
                text_parts = []
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        text_parts.append(part.get_content())
                return "\n".join(text_parts)
            else:
                return msg.get_content()

        except Exception as e:
            logger.error(f"Errore estrazione EML {file_path}: {e}")
            return f"[Errore lettura EML: {str(e)}]"

    def _extract_text_msg(self, file_path: str) -> str:
        """Estrae testo da file MSG (Outlook)."""
        try:
            import extract_msg

            msg = extract_msg.Message(file_path)
            text_parts = []

            if msg.subject:
                text_parts.append(f"Oggetto: {msg.subject}")
            if msg.body:
                text_parts.append(msg.body)

            msg.close()
            return "\n".join(text_parts)

        except ImportError:
            logger.warning("extract-msg non installato, impossibile estrarre testo da MSG")
            return "[extract-msg non disponibile - installare con: pip install extract-msg]"
        except Exception as e:
            logger.error(f"Errore estrazione MSG {file_path}: {e}")
            return f"[Errore lettura MSG: {str(e)}]"


# Singleton instance
_attachment_extractor = None


def get_attachment_extractor() -> AttachmentTextExtractor:
    """Recupera istanza singleton dell'estrattore."""
    global _attachment_extractor
    if _attachment_extractor is None:
        _attachment_extractor = AttachmentTextExtractor()
    return _attachment_extractor
